from io import BytesIO

from django.contrib import admin, messages
from django.db import transaction
from django.http import HttpResponse
from django.shortcuts import redirect, render
from django.urls import path, reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from .forms import ExcelImportForm
from .models import (
    Article,
    CardSet,
    Entry,
    Profile,
    SignupLog,
    Tag,
    Topic,
    TopicWord,
    TrainingResult,
)

EXCEL_HEADERS = ["Слово (англ.)", "Перевод", "Транскрипция", "Описание (англ.)"]
EXCEL_MAX_ROWS = 10000


def _cell(value, limit):
    if value is None:
        return ""
    return str(value).strip()[:limit]


def import_entries_from_excel(file, tags=(), update_existing=True, user=None):
    """
    Читает слова со второй строки первого листа: A — слово, B — перевод,
    C — транскрипция, D — описание. Возвращает словарь со счётчиками.
    """
    wb = load_workbook(file, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    stats = {"created": 0, "updated": 0, "skipped": 0}
    existing = {e.word.lower(): e for e in Entry.objects.all()}
    seen = set()

    with transaction.atomic():
        for i, row in enumerate(ws.iter_rows(min_row=2, max_col=4, values_only=True)):
            if i >= EXCEL_MAX_ROWS:
                break
            row = list(row) + [None] * (4 - len(row))
            word = _cell(row[0], 120)
            translation = _cell(row[1], 255)
            if not word or not translation or word.lower() in seen:
                stats["skipped"] += 1
                continue
            seen.add(word.lower())
            transcription = _cell(row[2], 120)
            description = _cell(row[3], 5000)

            entry = existing.get(word.lower())
            if entry is None:
                entry = Entry.objects.create(
                    word=word,
                    translation=translation,
                    transcription=transcription,
                    description=description,
                    created_by=user,
                )
                stats["created"] += 1
            elif update_existing:
                entry.translation = translation
                entry.transcription = transcription or entry.transcription
                entry.description = description or entry.description
                entry.save()
                stats["updated"] += 1
            else:
                stats["skipped"] += 1
            if tags:
                entry.tags.add(*tags)
    wb.close()
    return stats


def build_excel(entries):
    wb = Workbook()
    ws = wb.active
    ws.title = "Словарь"
    ws.append(EXCEL_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1E4FD8")
    for e in entries:
        ws.append([e.word, e.translation, e.transcription, e.description])
    for col, width in zip("ABCD", (28, 36, 22, 70)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


@admin.action(description="📤 Экспортировать выбранные в Excel")
def export_to_excel(modeladmin, request, queryset):
    data = build_excel(queryset.order_by("word"))
    response = HttpResponse(
        data,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    stamp = timezone.localtime().strftime("%Y-%m-%d")
    response["Content-Disposition"] = f'attachment; filename="eslovari-{stamp}.xlsx"'
    return response


@admin.register(Entry)
class EntryAdmin(admin.ModelAdmin):
    list_display = ("word", "translation", "transcription", "tag_list", "updated_at")
    list_filter = ("tags", "created_at")
    search_fields = ("word", "translation", "description")
    filter_horizontal = ("tags",)
    actions = [export_to_excel]
    list_per_page = 50
    change_list_template = "admin/dictionary/entry/change_list.html"
    fieldsets = (
        (None, {"fields": ("word", "translation", "transcription")}),
        ("Подробнее", {"fields": ("description", "example", "tags")}),
    )

    def get_queryset(self, request):
        return super().get_queryset(request).prefetch_related("tags")

    @admin.display(description="темы")
    def tag_list(self, obj):
        return ", ".join(t.name for t in obj.tags.all())

    def save_model(self, request, obj, form, change):
        if not change:
            obj.created_by = request.user
        super().save_model(request, obj, form, change)

    def get_urls(self):
        custom = [
            path(
                "import-excel/",
                self.admin_site.admin_view(self.import_excel),
                name="dictionary_entry_import",
            ),
            path(
                "excel-template/",
                self.admin_site.admin_view(self.excel_template),
                name="dictionary_entry_template",
            ),
        ]
        return custom + super().get_urls()

    def import_excel(self, request):
        if not self.has_add_permission(request):
            messages.error(request, "Нет прав на добавление слов.")
            return redirect("admin:dictionary_entry_changelist")
        form = ExcelImportForm(request.POST or None, request.FILES or None)
        if request.method == "POST" and form.is_valid():
            try:
                stats = import_entries_from_excel(
                    form.cleaned_data["file"],
                    tags=list(form.cleaned_data["tags"]),
                    update_existing=form.cleaned_data["update_existing"],
                    user=request.user,
                )
            except Exception:  # повреждённый или не-Excel файл
                form.add_error("file", "Не получилось прочитать файл. Это точно .xlsx?")
            else:
                messages.success(
                    request,
                    f"Импорт завершён: добавлено {stats['created']}, "
                    f"обновлено {stats['updated']}, пропущено {stats['skipped']}.",
                )
                return redirect("admin:dictionary_entry_changelist")
        context = {
            **self.admin_site.each_context(request),
            "opts": self.model._meta,
            "form": form,
            "title": "Импорт слов из Excel",
            "template_url": reverse("admin:dictionary_entry_template"),
        }
        return render(request, "admin/dictionary/entry/import_excel.html", context)

    def excel_template(self, request):
        sample = [
            Entry(word="deploy", translation="развёртывать (выкладывать)", transcription="[dɪˈplɔɪ]",
                  description="To release an application to a server so users can access it."),
        ]
        response = HttpResponse(
            build_excel(sample),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = 'attachment; filename="eslovari-shablon.xlsx"'
        return response


@admin.register(Tag)
class TagAdmin(admin.ModelAdmin):
    list_display = ("name", "slug", "description")
    search_fields = ("name",)
    prepopulated_fields = {"slug": ("name",)}


@admin.register(CardSet)
class CardSetAdmin(admin.ModelAdmin):
    list_display = ("title", "word_count", "is_published", "created_at")
    list_filter = ("is_published",)
    search_fields = ("title", "description")
    autocomplete_fields = ("entries",)
    prepopulated_fields = {"slug": ("title",)}

    @admin.display(description="слов")
    def word_count(self, obj):
        return obj.entries.count()


@admin.register(Article)
class ArticleAdmin(admin.ModelAdmin):
    list_display = ("title", "level", "reading_minutes", "is_published", "created_at")
    list_filter = ("level", "is_published", "tags")
    search_fields = ("title", "body")
    filter_horizontal = ("tags",)
    prepopulated_fields = {"slug": ("title",)}
    readonly_fields = ("reading_minutes",)


class TopicWordInline(admin.TabularInline):
    model = TopicWord
    extra = 0


@admin.register(Topic)
class TopicAdmin(admin.ModelAdmin):
    list_display = ("title", "owner", "created_at")
    search_fields = ("title", "owner__username")
    autocomplete_fields = ("entries",)
    inlines = [TopicWordInline]


@admin.register(Profile)
class ProfileAdmin(admin.ModelAdmin):
    list_display = ("user", "group", "current_login", "previous_login")
    search_fields = ("user__username", "user__email", "group")
    readonly_fields = ("current_login", "previous_login")
    exclude = ("favorite_entries", "favorite_sets", "favorite_tags")


@admin.register(TrainingResult)
class TrainingResultAdmin(admin.ModelAdmin):
    list_display = ("user", "mode", "source", "correct", "total", "created_at")
    list_filter = ("mode",)
    search_fields = ("user__username",)


@admin.register(SignupLog)
class SignupLogAdmin(admin.ModelAdmin):
    list_display = ("ip", "created_at")
    search_fields = ("ip",)
